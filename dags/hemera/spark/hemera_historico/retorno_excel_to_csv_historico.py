# Importing Libs
from minio import Minio
from io import BytesIO
import pandas as pd
from datetime import datetime, timedelta, timezone
from deltalake import write_deltalake
import os
import tempfile
import msoffcrypto
import time
import openpyxl

def retorno_excel_to_csv_historico (access_params=None, **kwargs):


    minio_raw = Minio(
    "api-raw.alpe.com.br",
    access_key = 'B7q0avvSIpSdyGPXWnEC',
    secret_key = 'PhMhRQSQ6YJU8fn2qKhDLM017cQPrlCz1YbM8IwU'
    )

    # Connection validation
    try:
        # Try to list the buckets
        buckets = minio_raw.list_buckets()
        
        # If the connection was successful, print the buckests
        print("Conexão bem-sucedida. Lista de buckets disponíveis:")
        for bucket in buckets:
            print(bucket.name)
        
    except Exception as e:
            # If the connection was failed, print the error message
            print(f"Erro ao conectar ao MinIO: {e}")
    

    # Configurações
    BUCKET_SOURCE = "hemera"
    BUCKET_DESTINATION = "hemera-vini"
    BASE_FOLDER = "retorno"
    YEARS = [2020, 2021, 2022, 2023, 2024, 2025]
    MONTHS = [f"{m:02}" for m in range(1, 13)]
    EXCEL_PASSWORD = "Alpe38"

    erro_count = 0
    arquivos_com_erro = []

    # Loop por anos e meses
    for year in YEARS:
        for month in MONTHS:
            folder_path = f"{BASE_FOLDER}/year={year}/month={month}/"
            print(f"🔍 Procurando arquivos em: {folder_path}")

            # Lista arquivos .xlsx na pasta
            objects = minio_raw.list_objects(BUCKET_SOURCE, prefix=folder_path, recursive=True)

            for obj in objects:
                file_key = obj.object_name
                if file_key.endswith(".xlsx"):
                    print(f"📥 Baixando arquivo: {file_key}")
                    
                    try:
                        # Baixa o arquivo como bytes
                        response = minio_raw.get_object(BUCKET_SOURCE, file_key)
                        raw_data = BytesIO(response.read())
                        decrypted_data = BytesIO()

                        try:
                            # Tenta descriptografar com a senha
                            office_file = msoffcrypto.OfficeFile(raw_data)
                            office_file.load_key(password=EXCEL_PASSWORD)
                            office_file.decrypt(decrypted_data)
                            decrypted_data.seek(0)
                            temp_file_data = decrypted_data.read()
                        except Exception:
                            # Se não estiver criptografado, usa raw_data direto
                            raw_data.seek(0)
                            temp_file_data = raw_data.read()

                        # Salva temporariamente para possível modificação via openpyxl
                        with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx") as temp_file:
                            temp_file.write(temp_file_data)
                            temp_file_path = temp_file.name

                        # Verifica se existe a aba 'Duplicatas' e ajusta se necessário
                        try:
                            workbook = openpyxl.load_workbook(temp_file_path)
                            if 'Duplicatas' in workbook.sheetnames:
                                aba_duplicatas = workbook['Duplicatas']
                                for sheet in workbook.sheetnames:
                                    if sheet != 'Duplicatas':
                                        del workbook[sheet]
                                aba_duplicatas.title = 'Sheet1'
                                workbook.save(temp_file_path)
                                print(f"📝 Mantida apenas a aba 'Duplicatas' em {file_key}")
                            else:
                                print(f"ℹ️ Nenhuma aba 'Duplicatas' encontrada em {file_key}, seguindo com conteúdo original.")
                        except Exception as e:
                            print(f"⚠️ Erro ao ajustar abas do Excel: {e}")

                        # Lê o arquivo ajustado
                        df = pd.read_excel(temp_file_path)

                        # Converte para CSV em memória
                        csv_buffer = BytesIO()
                        df.to_csv(csv_buffer, index=False)
                        csv_buffer.seek(0)

                        # Cria caminho do novo arquivo .csv
                        csv_key = file_key.replace(".xlsx", ".csv")

                        # Envia CSV para o novo bucket
                        minio_raw.put_object(
                            BUCKET_DESTINATION,
                            csv_key,
                            csv_buffer,
                            length=csv_buffer.getbuffer().nbytes,
                            content_type='text/csv'
                        )

                        print(f"✅ CSV enviado: {csv_key}")

                        # Remove o temporário
                        os.remove(temp_file_path)

                    except Exception as e:
                        erro_count += 1
                        arquivos_com_erro.append(file_key)
                        print(f"⚠️ Erro ao processar {file_key}: {e}")

    # Finaliza
    print("✔️ Processo finalizado.")
    print(f"❌ Total de arquivos com erro: {erro_count}")
    if arquivos_com_erro:
        print("🛑 Arquivos com erro:")
        for arq in arquivos_com_erro:
            print(f"   - {arq}")