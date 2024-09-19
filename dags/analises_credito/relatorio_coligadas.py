from datetime import datetime
import io
from airflow_dags_core.lib.MinioSaveIndex import MinioSaveIndex
from scripts.utils import extract_path_from_url, get_trino_connection, execute_query
from airflow.decorators import dag, task
from airflow.utils.dates import days_ago
from airflow_dags_core.lib.MinioWriteFile import MinioWriteFile
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment

default_args = {
    'owner': 'Rafael Leite',
    'retries': 1,
}

def create_xlsx(rows):
    wb = Workbook()
    ws = wb.active

    header_fill = PatternFill(start_color="782170", end_color="782170", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    header_alignment = Alignment(horizontal="center", vertical="center")

    row_fill_odd = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    row_fill_even = PatternFill(start_color="F6DEF4", end_color="F6DEF4", fill_type="solid")
    row_font = Font(color="000000")

    headers = ["CPF/CNPJ Sacado", "Nome Sacado", "Limite Atribuido", "Limite Utilizado", "Limite Disponivel"]
    ws.append(headers)

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment

    for idx, row in enumerate(rows, start=2):
        limite_utilizado = row[2] - row[3]
        ws.append([row[0], row[1], row[2], limite_utilizado, row[3]])

        for cell in ws[idx]:
            cell.font = row_font
            cell.fill = row_fill_odd if idx % 2 == 0 else row_fill_even

    for column in ws.columns:
        max_length = 0
        column = list(column)
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(cell.value)
            except TypeError:
                pass
        adjusted_width = max_length + 2
        ws.column_dimensions[column[0].column_letter].width = adjusted_width

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

def consulta_coligadas(cnpj, conn):
    sacados = None

    queryColigadas = f"""
        SELECT
            substring(documento, 2, length(documento)) AS cnpj_raiz
        FROM
            deltalaketrusted.serasa.coligada c
        WHERE
            id = (
            SELECT
                max(id)
            FROM
                deltalaketrusted.serasa.organizacoes o
            WHERE
                o.cnpj_raiz = '{cnpj[:8]}')
            AND c.identificacao_pessoa = 'J'
        """

    result = execute_query(conn, queryColigadas)

    if result:
        sacados = [row[0] for row in result]
    else:
        print(f"NÃO HÁ COLIGADAS PARA O SACADO {cnpj}")

    return sacados

@dag(
    default_args=default_args,
    start_date=days_ago(1),
    schedule_interval=None,
    description='DAG consulta o limite das empresas coligadas e envia o arquivo retornado para o MinIO',
    catchup=False,
)
def relatorio_coligadas():
    @task()
    def start_task():
        print("Operação iniciada")

    @task()
    def consulta_limites(**kwargs):
        conn = get_trino_connection()
        MinioIndex = MinioSaveIndex()
        try:
            issueKey = kwargs['dag_run'].conf.get('issue_key', 'default_value')
            cnpjSacado = kwargs['dag_run'].conf.get('payer_identification', 'default_value')
            if not cnpjSacado:
                raise ValueError("CNPJ não fornecido na execução da DAG.")

            sacados = consulta_coligadas(cnpjSacado, conn)

            if not sacados:
                return
            
            all_rows = []
            for cnpj in sacados:
                
                queryLimite = f"""
                    SELECT 
                        cpf_cnpj_sacado,
                        nome_sacado,
                        CASE WHEN categoria = 'SEGREGADO' THEN SUM(limite_atribuido) ELSE limite_atribuido END AS limite_atribuido,
                        CASE WHEN categoria = 'SEGREGADO' THEN SUM(limite_disponivel) ELSE limite_disponivel END AS limite_disponivel
                    FROM (
                        SELECT 
                            cpf_cnpj_sacado,
                            nome_sacado,
                            limite_atribuido,
                            limite_disponivel,
                            categoria,
                            pgid
                        FROM 
                            postgres.ccred_schema_prd_default.vw_limite_sacado_v3
                        WHERE 
                            cnpj_raiz = '{cnpj}'
                            AND cedente_principal = true
                        GROUP BY cpf_cnpj_sacado,
                                nome_sacado,
                                limite_atribuido,
                                limite_disponivel,
                                categoria,
                                pgid  
                    ) t1
                    GROUP BY cpf_cnpj_sacado,
                            nome_sacado,
                            limite_atribuido,
                            limite_disponivel,
                            categoria
                """
                
                rows = execute_query(conn, queryLimite)
                if rows:
                    all_rows.extend(rows)
                else:
                    print(f"Não há registro de limite para o CNPJ {cnpj}")
            
            if not all_rows:
                return print("Não há qualquer registro de limite dentre as empresas coligadas")
            
            xlsx_file = create_xlsx(all_rows)
            
            bucket = extract_path_from_url(kwargs['dag_run'].conf.get('minio_url', 'default_value'))
            
            current_time = datetime.now().strftime('%Y-%m-%dT%H:%M:%S.%f')
            
            fileName = f"coligadas_{cnpj}_{current_time}.xlsx"

            MinioWriteFile().write_file(file=xlsx_file, file_name= fileName, bucket=bucket)
            MinioIndex.save(key= issueKey, identification=cnpjSacado, path= bucket + f"/{fileName}", fileType= "COLIGADAS")
        finally:
            conn.close()

    @task()
    def end_task():
        print("Operação finalizada")

    start = start_task()
    limites = consulta_limites()
    end = end_task()

    start >> limites >> end 

coligadas_dag = relatorio_coligadas()